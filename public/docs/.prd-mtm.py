# Builds the MTM-edition PRD as HTML (localhost docs) and DOCX from one content tree.
import html as H, sys, datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DATE = '20 Oct 2026'
TITLE = 'BSC PLM: Product Requirements, MTM edition'
SUB = 'Made-to-measure end to end: the design loop, sourcing, the fabric master as Enigma holds it, purchase orders, the store, daily production and the factory floor. Workflows, user flows step by step, data sources, rules, build order.'

# content: list of ('h2', text) | ('p', text) | ('b', 'Bold lead.', 'rest') | ('table', headers, rows) | ('ul', items)
C = []
def h2(t): C.append(('h2', t))
def p(t): C.append(('p', t))
def b(lead, rest): C.append(('b', lead, rest))
def table(headers, rows): C.append(('table', headers, rows))
def ul(items): C.append(('ul', items))
def steps(rows):
    table(['#', 'Step · who', 'User flow', 'Data in (source)', 'System checks', 'Then'], rows)

# ------------------------------------------------------------------ 1
h2("1. Purpose and scope")
p("Requirements for the MTM edition of BSC PLM: Bombay Shirt Company's made-to-measure business end to end, as the MTM prototype shows it, screen by screen. Workflows run in the order a shirt happens: the order at the store, the day's production and the factory floor, then the cloth (fabric master, purchase orders, sourcing, the design loop, the swatch desk). RTW-only work is out and comes in a later phase; the full prototype stays its reference.")
p("BSC PLM sits beside Enigma, not inside it. Enigma stays the system of record for fabrics, vendors, customer orders, purchase orders and stock. BSC PLM reads what it needs over Enigma's API and keeps its own records for what Enigma does not hold: the design loop, vendor terms and holds, the swatch desk, production plans, scans, changed sheets, cases and conversation. Fields Enigma owns appear under Enigma's own names, so what is typed here is what goes there.")
b("In scope.", "The store's MTM work (bills, fit profiles, trial and express, sheet corrections, returns, credit, refunds); MTM daily production (queue, plan, pre-production desks, floor, changed sheets, extra fabric, ship-by moves), the factory merchandiser's floor and the Factories screen; the MTM fabric master as Enigma's fabric DB; purchase orders (auto or by hand) and the receiving line; the core weekly stock review; the sourcing desk and vendor master; the fabric requirement from ask to live online (five states); the swatch desk up to the handover to Yusuf; conversation and files on every record; the bell and names rule.")
b("Conventions.", "Dates are the record's own. Quantities are pieces unless the unit says metres. A state is where a record is, never the next step. Rules marked later are specified but not built in the prototype. Field names in italics are Enigma's.")

# ------------------------------------------------------------------ 2
h2("2. Users and roles")
p("Six roles, five with screens in BSC PLM (store managers and owners stay in RetailOS). One login per team; colleagues appear by name on what they wrote. Each role lands on the first screen it works in; the sidebar lists teams as pills.")
table(["Role", "People today", "Owns in BSC PLM", "Lands on", "Reads only"], [
 ["Design", "Kunal, Janak", "Every MTM fabric requirement; Push PO, More time or Reject on the vendor; designs and their push to Enigma; working, final and changed names; price per variant; go live online; pause and resume; promote to core; the fabric's Enigma fields", "Requirements", "Sourcing desk, swatch desk, purchase orders"],
 ["Sourcing", "Sreekanth, Srikanth", "A vendor on every ask; extension asks; the mills (Enigma's vendor master, Add and Edit Vendor); POs by hand; mill confirmation and remarks; the receiving line; the weekly core review; trims and packaging", "Fabric sourcing", "Requirements"],
 ["Production (MTM)", "Vishnu, Swati", "Confirm pre-confirmed orders; allocate to factories; plans; the pre-production overview; the factory floor; changed sheets; extra fabric; ship-by moves; Factories screen inputs", "MTM daily production", "Store orders"],
 ["Factory merchandiser", "Abhin M, Bhupendra, each unit's floor", "MTM floor: receive, out, request cloth, log the day, read changed sheets", "My factory", "Order sheet and job card"],
 ["Store", "Store managers (34 stores)", "No screen in BSC PLM, for managers or owners. Bills, profiles, sheet corrections, returns, credit notes and refund requests stay in RetailOS; everything additional is built in RetailOS only. The prototype shows them for demonstration", "Store", "Fabrics live today"],
 ["Owner", "Akshay, Parag", "Everything, read and write", "Any", "—"],
])
b("Warehouse and unit teams.", "No portal for Shravani and Yusuf in this edition; their work appears where it happens. Shravani books rolls on the PO's receiving line and approves the mill PO for bespoke and alteration extra fabric; Yusuf is the swatch desk handover. The CAD, cutting, fusing and compiling teams each work from their own desk. The warehouse returns as a role with its portal.")
b("Role rules.", "A button shows only for the role that owns the step; others see whose move it is. The factory merchandiser is scoped to one factory, the store to one store; a demo switch lifts both.")

# ------------------------------------------------------------------ 3
h2("3. Where the data comes from")
p("Enigma holds the transactions; BSC PLM reads them under Enigma's field names and records the decisions. Write-back is the last phase, only where access is granted.")
table(["Source", "BSC PLM reads", "BSC PLM writes (last phase, if granted)", "Recorded only", "Used by"], [
 ["Enigma · Fabric master (Edit Fabric Info)", "Fabric Name, Inventory Type, Category, Product Type, Fabric Code, On Hold; Vendor Information (Vendor Name, Vendor Code for fabric, Mill Country, Mill Description, Currency, Cost per Unit, TAT in hours, MOQ, Mills); Dynamic Attributes for shirts (Panna, Big-Checked, Thread Count, Fabric Ply, Wrinkle Resistance, Stretch, Consumed_In, Construction, Composition family, Washed Fabric, Wash shrinkage %, Washcare Description, Weight Class, Weave Description, Pattern, Fabric Color, Contrast Fabric, Fabric Piping, Is_Cutroll, Is_uniform, Weave, Composition); fabric image (URL, shown large on the pre-production desks)", "Fabric record edits; On Hold on pause, resume or supplier switch-off; the new fabric at Push PO", "Datasheet creation until write-back", "Fabric master, design loop, sourcing, store"],
 ["Enigma · Fabric inventory (Edit Inventory)", "Warehouse, Threshold, Max Inventory, Is Auto Replenishment, bins (code, metres, location on the rack, photo); physical, reserved, saleable", "Threshold and bin edits; adjustment with Adjustment Reason and Description", "—", "Fabric master, POs, store fabrics live today, pre-production desks (where the roll is)"],
 ["Enigma · Fabric transaction and reservation logs", "Every movement: date, order, PO, type, created by, warehouse, bin, reason, quantity; every reservation on punch", "—", "—", "Inventory & movements; rate of sale"],
 ["Enigma · Fabric PO", "Number, status, vendor, metres, created, overseer (the auto batch run), receiving line (ordered, received, pending, closed)", "Create at Push PO or Raise PO; Update Receiving Status", "PO number back on the record", "Design loop, sourcing, POs"],
 ["Enigma · Designs (Product Info)", "SKU, Product Type, Weight, Region, Stock Mode, Primary Fabric, Status, Sales Channel, Media; variant option and values with price, discounted price and SKU per variant", "Push per design; re-push on a price change", "—", "Design loop, store bill form"],
 ["Enigma · Orders and order sheet", "Order lines, states, status history, sheet fields, ship-by, QR", "Order-sheet corrections", "Punching the order", "Daily production, store, factory floor"],
 ["Enigma · Vendors", "Code, Name, Vendor Type, Address, Office Phone, Office Email, Active, Categories", "Add Vendor, Edit Vendor", "—", "Mills, sourcing"],
 ["Connected order software", "The order sheet XML per order", "—", "—", "CAD desk"],
 ["Shopify", "Sales by product code, customer profile", "—", "Listing, price, stock push", "Customer card, rate of sale"],
 ["Mill stock lists (email, WhatsApp)", "Article, metres, price", "—", "—", "Bespoke availability (store gate)"],
 ["Sheets, imported once", "Fabric ordering, Merchandise Master, revised names, swatch tracker, Trims Working", "—", "Read-only one month in parallel", "Seed and open items at go-live"],
 ["Typed in BSC PLM", "Requirements, vendor terms and holds, extensions, designs, names, swatch dates, plans, scans, desk steps, CAD files, corrections, cases, comments, files", "—", "—", "All workflows"],
])
b("Computed, never typed:", "saleable metres, weeks of cover, rate of sale, reorder point and suggested PO, headroom per factory, store credit balance, sheet version, every alert.")
b("Freshness.", "The fabric read is one call of about 3.2 MB with no change filter: polled nightly and on demand. Transaction reads need a code per call until a date range is confirmed. Shopify sales daily. Everything typed here is live.")

# ------------------------------------------------------------------ 4
h2("4. Workflow 1 — Store")
p("The MTM side of the store page: customers, orders, what I can sell, returns and alterations. Others can read it but not act. Store managers and owners have no screen in BSC PLM: they work in RetailOS, and everything additional for them is built there only. These flows fix what passes between RetailOS and BSC PLM; the prototype's store screens only demonstrate them.")
steps([
 ["1.1", "Look up the customer · store", "1. Customers: search by name or phone. The card shows home store, visits, spend, store credit with notes, measurements by fit profile, usual fabrics, past orders, open cases.\n2. Punch MTM order starts a bill with the last profile loaded.", "Shopify profile and spend; orders, cases, credit from BSC PLM", "none", "nothing written until a bill is punched"],
 ["1.2", "Punch a bill · store", "1. + Punch a bill, pick the customer, + Piece per item.\n2. Per MTM piece: fabric from fabrics live today (open for selling, with metres the store can still reserve), design, fit profile (a customer keeps several), who measured, quantity. The sheet opens pre-filled from the last order on that profile.\n3. Above the pieces: Trial shirt first for a new customer (about 25 days) or Urgent (7 to 8 days, ₹600); a comment for the factory; one for shipping.\n4. Apply store credit nets off up to the balance. Punch bill gives one bill number and one order number per piece with its promise date.", "fabric list and metres from Enigma; credit ledger", "only live fabrics; no sold-out bespoke fabric; credit capped at the balance", "metres reserved; pieces join production's queue after the freeze window"],
 ["1.3", "Follow the orders · store", "1. Orders lists every piece punched here with its stage, and Handed to customer.\n2. A ship-by moved by production reads “moved from 14 Oct, tell the customer” and lands in the bell.", "workflow 2", "none", "handover closes the piece"],
 ["1.4", "Correct the order sheet · store", "1. The sheet is Enigma's layout. Every correctable field is a dashed tile: click, type the correction and why, save.\n2. The heading reads V1 · as punched, then V2 · changed once, V3 · changed 2 times; the change log shows which version each change made.\n3. Once the piece is on a plan, measurements and styling lock and free-form remarks stop; only a logistics remark can change.", "typed", "edit only before production starts; logistics remark only after; nothing once made", "production sees the next version on Changed sheets"],
 ["1.5", "Return or alteration case · store manager", "1. Return or alter on a delivered piece shows whether it is inside 30 days (free alteration, credit note on return) or chargeable; reason, what the customer wants (Exchange, Credit note, Remake, Refund), remarks. Raise case gives a number and a 10-day promise.\n2. The ladder: Alter in store, Swap from the shelf, Send out, Refit, Remake, Take it back; each rung records its outcome; the case closes on the first that works.", "delivered pieces; typed", "only pieces delivered from this store", "take-back credit note: CN-<case>, MRP, valid one year, offered on the next bill"],
 ["1.6", "Refund approval · store, owner, finance", "1. Refund reads “requested, needs approval”. The owner approves from the bell; finance clicks Moved with the reference.", "amount from MRP", "approval before movement", "the ledger carries it"],
])
b("Alerts.", "Store: case past its promise; a fabric closed for selling; ship-by moved by production; refund awaiting approval (owner).")

# ------------------------------------------------------------------ 5
h2("5. Workflow 2 — MTM daily production and the factory floor")
p("Production runs its day from one screen: Allocate, Plans, Pre-production, Factory floor, Changed sheets, Extra fabric. The system proposes each plan from the factories' spare capacity; a person confirms. Every move of a piece is a scan. Daily output, aging and capacity inputs sit on the Factories screen.")
b("Stages, in order.", "Order (punched) → PP (waiting for a factory; frozen 2 hours for store changes) → PP · planned (on a plan) → pre-production at the unit: CAD → Cutting → Fusing → Compiling → In production (scanned out) → In factory (aging after 5 days) → Out factory (QC passed) → RTS (back at the unit) → Fitting / shipping → Delivered. Cancelled is possible before cutting.")
steps([
 ["2.1", "The queue fills · system", "1. A piece joins Allocate once its freeze (2 hours, a setting) has passed; pieces still frozen sit in a grey strip with the time they join.\n2. Urgent first, then oldest; trial and express keep their shorter promise; bespoke shows “waiting for cloth” until its parcel is in.", "orders punched in Enigma", "not allocatable inside the freeze; bespoke not before its parcel", "—"],
 ["2.1a", "Pre-confirmed orders · store, production", "1. An order taken but not confirmed (measurement not done, on hold, payment pending) is pre-confirmed, with its reason and a note.\n2. Allocate filters: Exclude pre-confirmed (the queue, default) and Pre-confirmed (oldest first, days pre-confirmed, Confirm).\n3. RetailOS gets the store's pre-confirmed list: order, customer, fabric, reason and note, days pre-confirmed (red past 7), oldest first, a reason filter and Confirm. The prototype store page demonstrates it.\n4. Confirm order, with an optional remark, moves it into the queue.", "the store; typed", "never planned to a factory until confirmed", "confirmed orders join in punch order"],
 ["2.2", "The system proposes · system", "1. One card per MTM factory: feed target, pieces on the floor and their age, in transit, headroom today, load in days, number proposed.\n2. Corporate to the corporate factory; urgent to the factory with an urgent reserve if it has room, else the nearest with room; the rest to the most headroom, oldest first.", "Factories inputs; day logs", "headroom = daily capacity − in transit − pieces beyond two days of work", "—"],
 ["2.3", "The planner confirms · production", "1. Tick pieces (all, oldest 20/30/50, or the proposal) and pick the factory; the card shows headroom after the plan and “over by N” in red.\n2. Submit N to <factory>; over headroom it reads Send anyway and records the override.", "typed", "over headroom refused unless forced; urgent sets ship-by = punch + 2 days", "a plan number; pieces read PP · planned; the factory sees “on its way”"],
 ["2.3a", "Open a plan · production", "1. Plans lists each plan: number, factory, created, pieces, a pill per stage with a count, open or closed.\n2. A plan shows its facts (factory, pieces, metres, fabrics, first ship-by, urgent count, any over-capacity send and by whom), a strip of where its pieces are (planned → in production → in factory → out factory → RTS → dispatched), the scans that apply now, the cut list by fabric, and every piece with order, customer, store, fabric, metres, ship-by, stage and its scan dates.\n3. A piece opens its order sheet with the full timeline.", "the plan; scans", "none", "dates are the scans; nothing typed"],
 ["2.4", "Pre-production desks · CAD, cutting, fusing, compiling teams", "1. After a plan is submitted every order passes four desks, each with its own screen: CAD (the order sheet XML becomes a CAD file and a print), Cutting (cloth cut to the CAD, print attached), Fusing (collar, cuff and placket interlining), Compiling (bundle checked and approved).\n2. Tabs: Waiting for us, Taken up, Blocked, Done. Take up puts the person's name on the orders; Mark done, with an optional remark, hands them to the next desk. An order that cannot move (XML missing a measurement, short roll, interlining out) is marked blocked with a reason until unblocked. Urgent sits on top. Filters: Urgent; search by order, customer or fabric; factory; plan. Each row shows the desks checklist (done, in progress or blocked, with who).\n3. Opening an order, on any desk, shows: the desks checklist with who took up and finished each step, plus Take up, Mark done or Mark blocked; the fabric image from the backend, large (largest on cutting), metres to cut, and the bin from the backend (code, place on the rack map, photo), contrast fabric likewise; the CAD (XML and file); the order sheet (plan, factory, ship-by, styling, measurements, remarks).\n4. CAD sees each order's XML (View XML, copy, download), generated by the connected order software, and uploads the CAD file it made; CAD cannot be marked done without it, and the file stays on the order for cutting.\n5. Production's Pre-production tab (Swati) shows each team's queue: ready for it, blocked (with reason), work in progress, hand-offs completed in 7 days, and orders ready to send. A number opens those orders with their four steps and any block.", "the plan; XML from the connected order software; fabric image, bin location and photo from the backend; CAD file uploaded; typed", "desks in order; a block needs a reason; CAD needs its file; leaves for the factory only once compiled", "who did each step and when is on the order"],
 ["2.4a", "Contrast fabric · compiling team", "1. A piece can carry a little of another inventory fabric (pocket, collar band, cuff lining), listed on the order with fabric and metres; compiling can add a line.\n2. Approve & mark compiled deducts the contrast metres from that fabric's physical stock, with a transaction line.", "the order sheet; typed at compiling", "inventory fabric only; metres per line", "physical falls by the contrast metres; saleable follows"],
 ["2.4b", "Scan out · unit", "1. Each plan has a cut list by fabric and the order sheets. Scan out sends compiled orders to the factory.", "the plan; consumption from Enigma", "orders not yet compiled are held back", "In production; the main fabric's metres leave physical and reserved stock"],
 ["2.5", "The factory scans · merchandiser", "1. My factory: on their way, on the floor, ready. Received per piece or selection; Out of the factory when done; Log the day with pieces out and a note.\n2. Production's Factory floor shows the same per factory: order, customer, fabric, metres, tag, ship-by, stage, reached (or sent) date, days in factory (red past 5). Urgent sits on top, with an Urgent filter so the floor works them first. A row opens the order sheet.", "scans; typed counts", "each piece must be at the previous stage", "Factories screen shows the factory's count beside the scanned, amber when they differ"],
 ["2.6", "A change after production starts · store, production, factory", "1. Once on a plan, no measurement or styling change reaches the factory: stray remarks get missed and leave the factory blamed for a change it could not make.\n2. A logistics remark shows on the order for shipping; nothing is re-issued to the factory.\n3. Changes made before planning appear on Changed sheets; production clicks Acknowledged once it has acted.", "typed", "free edits until on a plan; logistics remark only after; nothing once made", "the factory never gets a changed sheet mid-production"],
 ["2.7", "Move a ship-by · production", "1. One piece: Factory delayed, move ship-by with a date and why. A whole floor: factory, days, one reason.", "typed", "open pieces only", "original date kept; the store reads “moved from…, tell the customer”; the portal marks “moved”"],
 ["2.8", "Extra fabric · factory, alteration unit, production, Shravani", "1. The merchandiser clicks Ask for extra cloth: cloth, panel (front 0.8 m, both fronts 0.9, sleeve 0.8, both sleeves 0.9, collar 0.3, cuff 0.3), metres, reason (factory or PP error), note. The alteration and hemming unit asks too (part replacement, CCP collars and cuffs for stores).\n2. Swati sees each ask with the shelf before and after, filtered: Normal orders, Bespoke orders, Alteration / hemming. Cloth on the shelf: Approve or Decline with a note; approval deducts the metres.\n3. Bespoke, alteration, or a shortfall the shelf cannot cover goes to Shravani in one click. An alteration the shelf covers is issued and deducted from physical stock on her approval; otherwise she approves, the PO is generated, and when the cloth is in the ask returns to Swati to approve.\n4. Bespoke is never inventory: nothing booked in on arrival, nothing deducted on approval.", "typed; PO number from the mill PO", "needs a fabric and a factory; later capped at the panel allowance", "approval deducts at once; factory errors count on the vendor, PP errors on the consumption rule"],
 ["2.9", "Round, dispatch, delivery · unit, shipping, store", "1. RTS at the unit, packed, shipped; the store hands over (1.3).", "scans", "none", "delivered"],
])
b("Factories screen.", "Capacity and planning inputs per factory (tailors, nominal and measured output, feed target, distance, corporate, urgent reserve, grade), Daily output (the WhatsApp table, keyed beside scanned), Aging (past ship-by by factory, stage, days), On-time performance.")
b("Alerts.", "Production: past ship-by; changed sheets pending; extra fabric to decide. Factory: unread changes; pieces to receive; a ship-by moved.")

# ------------------------------------------------------------------ 6
h2("6. Workflow 3 — The MTM fabric master, as Enigma holds it")
p("One list and one record: the list is Enigma's fabric inventory with its columns; the record is Enigma's fabric DB, its screens shown as tabs.")
steps([
 ["3.1", "Fabric inventory · design, sourcing", "1. What sells → Fabric inventory lists every MTM fabric: code, name, category, inventory type, saleable, physical, reserved, threshold, status.\n2. Filters: swatch book strip, Core / Non-core, On the shelf, warehouse; export.\n3. The sell-status cell carries one action for its state: Pause selling…, Resume, Switch back on, Switch off · supplier.", "Enigma fabric master and inventory", "one action per state", "pause and switch-off set On Hold"],
 ["3.2", "The fabric page · everyone", "1. A fabric opens as its own page (#/fabric/<id>): a facts strip (on the shelf, reserved, on order, saleable, threshold, cost per metre, weekly sales, status) and five tabs.\n2. Overview: bucket, promote or demote, core supply block, family, vendor, book; 12 weeks of sales; in stores.\n3. Purchase orders: how it is reordered (Enigma or a person), the reorder point and whether it is under; every PO with its source pill and receiving line.\n4. Fabric info · Enigma: Edit Fabric Info, Vendor Information, Dynamic attributes for shirts, each editable.\n5. Inventory & movements: Edit Inventory (Warehouse, Threshold, Max Inventory, Is Auto Replenishment, bins with code, metres, rack location and photo), Adjust stock with reason and description, the transactions and reservations logs.\n6. Files & conversation.\n7. Rename keeps the old name in history; the code never changes.", "Enigma; GRNs; cuts; typed adjustments", "rename needs a new name; adjustments are signed metres with a reason", "every edit is what goes to Enigma; the same record opens as a drawer from the list"],
])

# ------------------------------------------------------------------ 7
h2("7. Workflow 4 — Purchase orders and supply")
p("A PO is raised one of two ways, said plainly on screen. Enigma · auto: its batch run orders a domestic core fabric once saleable falls under the threshold. By hand: sourcing for international core (Enigma never orders it); design for a new fabric (workflow 6).")
steps([
 ["4.1", "Purchase orders · sourcing", "1. Purchase orders & supply → Purchase orders lists every open PO with a pill: Enigma · auto, by hand · design, by hand · sourcing.\n2. Filters: All open, Enigma · auto, By hand, Needs a PO by hand, Mill not confirmed, Confirmed.\n3. Needs a PO by hand lists international core under its reorder point with nothing on order, each with Raise PO · N m.\n4. When the mill replies: Mill confirmed… with date and remark; later Update date / remark.", "Enigma fabric POs; computed reorder points", "Raise PO only under the point with no open PO", "an unconfirmed auto-PO is a chase after 2 days; the receiving line (6.5) books rolls"],
 ["4.2", "Core · weekly stock review · sourcing", "1. One table: fabric, mill, saleable, reorder state (above or under, PO if any, Raise PO for international), the mill's last word and when, Review (In stock, Delayed, Run out, back-on date, note), and the selling cell.\n2. Filters: All, Review due, Under the reorder point, PO by hand, Switched off.", "typed", "review due if never asked or 7+ days ago", "delayed or run out goes to design to switch off; back in stock, they switch on"],
 ["4.3", "Trims & packaging · sourcing", "1. The trims master: stock, threshold, incoming, cost.", "Trims Working sheet; receipts", "none", "—"],
])
b("The rule.", "Saleable = physical − reserved − floor. Rate = shirts a week over 8 weeks. Metres a week = rate × 2.15. Lead-time demand = metres a week × lead days ÷ 7. Reorder point = the larger of lead-time demand and the vendor's MOQ; under it when saleable + on order ≤ the point. PO = the MOQ. Never for bespoke. Promotion to core lifts the floor to 40 m if under 20.")

# ------------------------------------------------------------------ 8
h2("8. Workflow 5 — Sourcing desk and the mills")
p("Sourcing sees only its own moves: add a vendor to a new ask, or answer an extension. Two tabs that never overlap: Your move (state 1 and extensions) and Waiting on design (state 2). From state 3 a record is on Purchase orders instead.")
steps([
 ["5.1", "Your move · sourcing", "1. Each row is one ask: what design asked for (name, why, book, month, attachments) and the vendor (empty, or the extension asked).\n2. Add vendor (6.2), or Extend / Drop vendor (6.4).", "requirements", "none", "moves to Waiting on design"],
 ["5.2", "Waiting on design · sourcing", "1. State 2 rows with the vendor card as design sees it and how long design has had it.", "requirements", "none", "—"],
 ["5.3", "Find a fabric · sourcing", "1. Fabric vendors → Find a fabric: search cloth, code, article, composition, book or mill; filter by mill, origin, bucket, book; sort by name, cheapest, fastest, least cover.\n2. Each row: fabric and code, composition and panna, vendor, TAT · MOQ, cost in the mill's currency, buyable today yes or no (hover for why), and how many other mills do it, with a compare panel.\n3. A row opens the fabric page.", "fabric master; mills", "none", "—"],
 ["5.4", "The mills · sourcing", "1. The mills is Enigma's vendor list: Code, Name, Vendor Type, Address, Office Phone, Office Email, Categories, Active, plus TAT · MOQ and fabric count.\n2. + Add Vendor is Enigma's form plus Vendor Information (Mill Country, Mill Description, Currency, TAT in hours, MOQ); saving opens the mill.\n3. A mill opens on Vendor · as in Enigma with Edit Vendor, then its fabrics, performance, blocks, files and conversation.", "Enigma vendors; typed", "name required", "goes to Enigma at write-back"],
])

# ------------------------------------------------------------------ 9
h2("9. Workflow 6 — Design: an MTM fabric, requirement to live online")
p("A new fabric, shade or construction variant runs on one record, always in one of five states, shown in one column: 1 Requirement requested (sourcing adds a vendor) · 2 Vendor added (design punches the PO) · 3 PO punched (cloth on its way) · 4 Fabric in (design configures Enigma; design codes can be added at any stage) · 5 Live online. The record opens on four sections in that order; PO punched is a state, not work, so it has no section.")
b("Trigger.", "A gap in a swatch book, a customer ask, or a variant worth making. Done when live online with its designs in Enigma. Stores follow the swatch desk and do not gate going live online.")
steps([
 ["6.1", "Raise the requirement · design", "1. + New requirement: a new fabric or shade, or a variant of an existing one (picked from the master).\n2. Working name, why BSC needs it, swatch book, month wanted, a sketch, swatch photo or pattern file, and whether it will be washed (shrinkage is known only once the cloth arrives). Samples are an optional checklist: Sample & PAT file to a factory, PP sample approved, Shipment sample approved, Approval posted to every department; one click each with an optional remark, a remark required to send a sample back; remarks land in the conversation under their step.\n3. Raise: state 1 · Requirement requested.", "typed; base fabric from the master", "name and why required; a variant's base must be inventory over 20 m, not bespoke, not paused or on hold", "the ask lands on sourcing's Your move"],
 ["6.2", "Sourcing adds the vendor · sourcing", "1. On Your move, Add vendor: mill, article, cloth name, price per metre, MOQ, lead days, hold-until (today plus the mill's usual hold), composition, panna, book, note.\n2. Save vendor: state 2 · Vendor added; the row moves to Waiting on design.", "typed; mill from the vendor master", "mill, price and cloth name required; one live vendor per ask, a second replaces the first, which is kept as replaced", "design is reminded; urgent at 5 days of hold left"],
 ["6.3", "Design answers on the vendor card · design", "1. The card shows terms (price, MOQ, lead, hold-until, article, composition, panna) and the first buy.\n2. Metres default to the MOQ; the amount follows as typed.\n3. Push PO, More time, or Reject with a reason. No sample gate here.", "metres typed; price from the card", "Push PO refused on a lapsed hold until extended; Reject needs a reason", "Push PO: the PO is punched (Enigma's number, or a stand-in FPO until write-back), the fabric joins the master as non-core, on hold; state 3 · PO punched. More time: an extension ask goes to sourcing. Reject: back to state 1"],
 ["6.4", "Extend or drop · sourcing", "1. An extension appears on Your move: type the new hold date and Extend, or Drop vendor if the mill will not hold.", "typed", "new date must be later (later)", "back to 6.3"],
 ["6.5", "The PO and its receiving line · sourcing, warehouse", "1. Section 2 shows the PO: number, metres, amount, raised on, mill confirmation.\n2. Enigma's receiving line below it: Receiving Status, Ordered, Total Received, Pending, a Received Qty box, Close, Update Receiving Status.\n3. Rolls land in parts; each update books metres and raises a GRN; Close with metres pending cancels the balance.\n4. Washed fabric: a cut is washed on arrival and the receiver enters the shrinkage %. The GRN keeps metres received; the shelf books usable metres (100 m at 5% books 95 m); the difference is a WASH transaction.\n5. The same line is on Purchase orders.", "PO from 6.3; metres counted; washed from the requirement; shrinkage at receiving", "received never above pending; shrinkage only for fabric marked washed, required before booking", "shelf up by usable metres, incoming down by what arrived; received or closed moves to state 4 · Fabric in"],
 ["6.6", "Designs and stickers · design", "1. Design codes can be added at any stage. In Designs & stickers, its own section, + Add a design (the form stays closed until clicked) fills Enigma's Product Info: SKU, Product Type, Weight, Region, Stock Mode, Primary Fabric (this cloth), Status, Sales Channel, Media.\n2. A variant option (Size, Fit, anything) with free-text values; per value a price, discounted price and SKU (default code-value).\n3. Description and buzzwords for the customer.\n4. Save mints a product code (BOM-SHI-FM/CS-nnnn); Push to Enigma sends it and the row reads “in Enigma · synced <date>”.\n5. Stickers sit under the designs: one per design (fabric code, design name, product code, swatch book, MRP), printed 3 × 7 or one page per design; no per-store copy count.", "typed", "a design needs a name and one variant with a price; price only on variants", "listed on Design → Designs (every product code, searchable, in Enigma or not) and in Search; offered on the store's bill form; a price change re-pushes and stamps the sync date"],
 ["6.7", "Configured in Enigma · design", "1. Set the fabric up in Enigma (template, base and buffer consumption) and click Configured in Enigma with those values. Big-Checked and composition family are fabric attributes, edited on the fabric record.", "typed; later read back", "template and base metres required", "the sheet's consumption reads base + buffer; Big-Checked adds matching"],
 ["6.8", "Name and price · design", "1. Working name, final name (settable any time) and every earlier name with who, when and why.\n2. Price is per variant (6.6).", "typed", "none", "website, sticker and Enigma read the final name; a change after a sticker flags “out of date · reprint”"],
 ["6.9", "Live online · design", "1. Section 4 lists the online gates (photos, video, post-production, product page) and Go live online, in any order. Stores are not a gate; story and live in stores are hidden in this edition.", "typed", "blocked until composition and a description are on the record; other gates advisory, skips kept", "state 5 · Live online; off hold for the website"],
])
b("Cases.", "New shade of a known fabric: no sampling. Re-source from another mill: a new ask. Paused before live: hidden from website and order form; the record stays until design closes it. Core replenishment and bespoke are never requirements.")
b("Alerts.", "Design: vendor added awaiting Push PO (urgent at 5 days of hold left); fabric in with no design; a design not pushed. Sourcing: a requirement with no vendor; an extension asked.")

# ------------------------------------------------------------------ 10
h2("10. Workflow 7 — Swatch desk, phase one")
p("Here the desk runs from cloth in to the handover to Yusuf; cut, sticker, dispatch, photos and stores told follow with the warehouse portal.")
steps([
 ["7.1", "Cloth in → to Yusuf · warehouse", "1. Every new fabric with received metres, each with Handed to Yusuf (tick Priority-1 if urgent).\n2. The row keeps dates, who, files and conversation.", "the receiving line (6.5)", "steps in order", "design sees “with Yusuf”; Priority-1 reminds until done"],
])
b("Phase two.", "Yusuf cuts swatches for every store, prints the sticker (fabric code, design name with article, book, price), dispatches one parcel per store and takes photos; design marks stores told; a name or price change after printing flags a reprint.")

# ------------------------------------------------------------------ 11
h2("11. Workflow 8 — Conversation and files on every record")
p("Nothing is emailed or kept in WhatsApp: the record is the thread. Every record (order, customer, fabric, requirement, mill, return case) has a conversation and files at its foot, open to anyone who can open the record. A mention puts a line on that person's bell until they mark it dealt with.")
steps([
 ["8.1", "Post · anyone who can open the record", "1. Conversation · N lists comments in order with who, role, when and step.\n2. Pick the step (a requirement: General, Fabric, PP sample, Shipment sample, Approval, Enigma; an order: General, Change from store, Fabric, Alteration; a customer: General, Fit, Preferences, Complaint; the swatch desk: Cut & packed, Dispatch, Marketing · go live, General), type, attach files if any, Post.\n3. Filter by step.", "typed; files", "text required", "on the record for everyone; the toast says who will see it"],
 ["8.2", "Mention · writer", "1. @ and a name or team (@kunal, @sourcing, @design), resolved against the people and roles list.\n2. A store writing on an order mentions whoever holds the piece by default: production for MTM in PP or at a factory.", "typed", "unknown names stay plain text", "each person, or every member of a team, gets a line under Messages"],
 ["8.3", "Deal with a mention · the person mentioned", "1. Waiting on you → Messages: newest first, by day, with the record (name, code), step, text and writer. Filters: all, this week, earlier.\n2. The text opens the record at the conversation.\n3. Mark resolved keeps the comment with who resolved it and when; Resolve all clears the list; a resolved mention can be put back.", "threads", "done when dealt with, not when read", "the bell count drops; the writer sees who resolved it"],
 ["8.4", "Files · anyone who can open the record", "1. Files · N with a label per record kind. Requirement: PAT file, CAD / pattern, tech pack, sketch or reference, sample photo, green slip, approval note, swatch image. Fabric: swatch image, test report, mill stock sheet, vendor quote, block confirmation, datasheet. Order: order sheet, alteration note, photo, measurement sheet. Customer: fit sample photo, reference image, measurement sheet, old bill. Mill: agreement, rate card, capacity sheet, audit report. Return: photo, invoice, alteration note. Swatch desk: cut swatch photos, dispatch photo, sticker page, marketing lines, reference.\n2. Upload, pick the label and optionally the step; images preview inline; a file on a comment is labelled Attachment.", "files", "25 MB per file", "stays with the record; sticker and website read the swatch image and marketing lines here"],
 ["8.5", "System posts · system", "1. Some actions write their own line so the record tells its story: an approval is posted once under Approval for every role; a rename, pause, ship-by move or refund approval leaves who and why.", "actions", "none", "audit and conversation are one thread"],
])
b("Rules.", "One thread per record, in order, never deleted, no private comments. A team mention reaches every member. Resolving is per person and kept on the comment. Nothing goes by email or WhatsApp.")
b("Alerts.", "Messages on the bell: every unresolved mention of me or my team. The bell counts things to do plus unresolved mentions.")

# ------------------------------------------------------------------ 14
h2("12. Business rules and validations")
table(["Area", "Rule"], [
 ["Requirements", "Name and why required. A variant's base must be inventory, over 20 m, not bespoke, not paused or on hold. One of five states, derived, never typed."],
 ["Vendors and holds", "One live vendor per ask; a second replaces the first, kept as replaced. Push PO refused on a lapsed hold until extended. Reject needs a reason. Extend must be later (later). PO amount = typed metres × the vendor's price."],
 ["Purchase orders", "Enigma-raised is source auto; person-raised is manual with who. Received never exceeds pending; Close with metres pending cancels the balance as closed short. Every receiving update raises a GRN. Needs a PO by hand is only international core under its point with nothing on order."],
 ["Designs", "A name and one priced variant. Price per variant only. A price change on a design in Enigma re-pushes it. The product code is minted once."],
 ["Names", "A rename appends the old name with who, when and why; the code never changes."],
 ["Fabric master", "Edits saved under Enigma's field names. An adjustment is signed metres with a reason and description. Threshold is the reorder floor. One sell action per state."],
 ["Go live online", "Blocked until composition and a description are on the record; other gates advisory, skips kept. Stores are not a gate."],
 ["Swatch desk", "Steps in order; phase one ends at the handover to Yusuf."],
 ["Extra fabric", "Approval from the shelf deducts from physical stock at once, alteration cloth included. Bespoke, an alteration ask, or a shortfall goes to Shravani in one click; an alteration the shelf covers is issued and deducted on her approval; otherwise she approves, the PO is generated, and it returns to production when the cloth is in. Bespoke is never inventory: nothing booked in or deducted. A fuller tool for Swati comes later."],
 ["Washed fabric", "Washed is set on the requirement; shrinkage % is entered at receiving once a cut is washed. Only washed fabric shrinks: stock books received × (1 − shrinkage); the GRN keeps metres received."],
 ["Pre-confirmed orders", "Never planned to a factory; shown only under the Pre-confirmed filter until confirmed."],
 ["Pre-production", "Every planned order passes CAD, cutting, fusing and compiling in order, each step recording who took it up, who finished and when. A block needs a reason. CAD needs its file. Scan-out is refused until compiled."],
 ["Contrast fabric", "Inventory fabric in small amounts; compiling approval deducts it from physical stock. Reserved is the main fabric held at punch; saleable = physical − reserved − floor."],
 ["MTM production", "A piece joins the queue after the freeze. Over headroom is refused unless forced, and forcing is recorded with the name. Urgent sets ship-by to punch + 2 days. Scan-out deducts from stock and reserved. Bespoke is never planned before its parcel. A ship-by move needs a reason."],
 ["Sheet changes", "Before production (order or PP) store and production can correct the sheet, each correction the next version (V2, V3). On a plan, measurements and styling lock and free-form remarks stop; only a logistics remark can change."],
 ["Store", "Only live fabrics can be punched; no sold-out bespoke. Credit never above the balance; credits expire after a year. Cases only on pieces delivered from this store; within 30 days free alteration and credit note, after that chargeable. A full refund needs approval before finance moves it."],
 ["Replenishment", "Reorder point = the larger of lead-time demand and MOQ; PO = the MOQ. Never for bespoke. Promotion to core lifts the floor to 40 m if under 20. A fabric the mill reports out goes on hold at once."],
 ["Roles", "Buttons only for the owning role. Store and factory roles are scoped to one store or factory. Owner sees and may do everything."],
 ["Conversation and files", "One thread per record, never deleted, readable by everyone who can open it. A mention stays on the bell until resolved; resolving is kept. Files carry a label per record kind; 25 MB each. Nothing is emailed."],
])

# ------------------------------------------------------------------ 15
h2("13. Non-functional requirements and the Enigma contract")
table(["Area", "Requirement"], [
 ["Architecture", "A separate application beside Enigma: reads over its API, its own database for its records, write-back behind per-module flags in the last phase, nothing built in Enigma's codebase. One codebase builds both editions; MTM is a build flag hiding RTW screens and the warehouse portal."],
 ["Enigma reads", "Fabric master (one 3.2 MB call, no change filter) with the fabric image: nightly and on demand. Inventory per warehouse with each bin's code, rack location and photo. Transaction and reservation logs per code until a date range is confirmed. Fabric POs with receiving lines. Orders and order sheets. Vendors. Designs and variants."],
 ["Enigma writes (last phase, if granted)", "Fabric info, attribute and inventory edits; On Hold on pause, resume or switch-off; the fabric PO at Push PO and Raise PO; Update Receiving Status; Add and Edit Vendor; design push with variants; order-sheet corrections. Until then each is keyed in Enigma and the number recorded here."],
 ["Other reads", "Shopify sales by product code daily, and the customer profile. The order sheet XML from the connected order software. Mill lists pasted. Sheets imported once at go-live, idempotent, every skipped row logged."],
 ["Identity", "One login per person, one role, store or factory scope where it applies; an administrator sees everything."],
 ["Concurrency", "Server-side records with optimistic versioning; a stale write is refused and the current state shown; Enigma's reservation decides when two stores punch the last metres."],
 ["Audit", "Every write appends who, when and what. Order sheets keep versions. Nothing hard-deleted. History exportable."],
 ["Performance", "Server-side filtering for the fabric master (about 8,600 rows) and orders (about 40 a day). The bell recomputed under a second on every write."],
 ["Availability", "Stores work 10:00 to 22:00 IST, seven days; maintenance outside those hours."],
 ["Devices", "Desktop first; a narrow layout with large buttons for the returns bench and the swatch desk."],
 ["Files", "Object storage, 25 MB per file, images previewed inline, label per record kind."],
 ["Environments", "Development; staging reading Enigma read-only against production data; production. Enigma has no sandbox, so writes are flagged per module per environment."],
 ["Backups · localisation", "Daily, kept 90 days. English; dates as 12 Sep 26; rupees with Indian grouping; inches for measurements, metres for cloth."],
 ["Parallel run", "Each phase keeps the sheet it replaces read-only for one month, with a weekly comparison with the owning team before switch-off."],
])


# ===================================================================== HTML · devx Doctrine
import re
def nd(t):
    """The doctrine allows no em dashes: a dash between clauses becomes a comma, an empty cell says None."""
    t=t.strip()
    if t in ('—', '-'): return 'None'
    t=re.sub(r'\s+—\s+', ', ', t)
    t=t.replace('—', ', ')
    return t
def split_head(t):
    """'4. Workflow 1 — Store' -> ('04', 'Workflow 1', 'Store'); '12. Business rules' -> ('12', None, 'Business rules')"""
    m=re.match(r'(\d+)\.\s+(.*)$', t); num, rest = (m.group(1).zfill(2), m.group(2)) if m else ('', t)
    m2=re.match(r'(Workflow \d+)\s+—\s+(.*)$', rest)
    return (num, m2.group(1), m2.group(2)) if m2 else (num, None, rest)
def esc(t): return H.escape(t)
def para_html(t): return '<br>'.join(esc(nd(x)) for x in t.split('\n'))

FONTS='<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family=Inter+Tight:wght@300;400;500;600;700&family=Source+Serif+4:ital,opsz,wght@0,8..60,400;0,8..60,500;1,8..60,400&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet" />'
css = """
:root{--ink:#0A0A0A;--ink-2:#1A1A1A;--muted:#5C6066;--muted-2:#8A8F96;--rule:#E5E5E5;--rule-2:#F0F0F0;--paper:#FFFFFF;--paper-2:#FAFAF8;--paper-3:#F4F4F1;--accent:#1E6FFF;--accent-soft:#E8F0FF;--warn:#C0392B;--ok:#0A7C53;
--f-display:'Inter Tight',-apple-system,BlinkMacSystemFont,sans-serif;--f-serif:'Source Serif 4',Georgia,serif;--f-mono:'JetBrains Mono','SFMono-Regular',monospace}
*{box-sizing:border-box}
html{background:var(--paper)}
body{margin:0;font-family:var(--f-display);color:var(--ink);background:var(--paper);font-feature-settings:"ss01","cv11","tnum";-webkit-font-smoothing:antialiased;text-rendering:optimizeLegibility;font-size:15px;line-height:1.5}
.page{max-width:960px;margin:0 auto;background:var(--paper);padding:64px 32px 96px}
.eyebrow{font-family:var(--f-mono);font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);margin:0 0 12px}
.eyebrow b{color:var(--accent);font-weight:500}
h1{font-family:var(--f-display);font-weight:500;font-size:48px;line-height:1.05;letter-spacing:-.025em;margin:0 0 24px;color:var(--ink)}
h1 em,h2 em{font-family:var(--f-serif);font-style:italic;font-weight:400;color:var(--accent);letter-spacing:-.01em}
.lede{font-size:17px;color:var(--ink-2);max-width:720px;margin:0 0 32px}
.pill{display:inline-block;font-family:var(--f-mono);font-size:10px;letter-spacing:.14em;text-transform:uppercase;border:1px solid var(--ink);border-radius:999px;padding:4px 12px;margin-left:12px;vertical-align:middle;color:var(--ink)}
.meta{font-family:var(--f-mono);font-size:11px;letter-spacing:.08em;color:var(--muted);text-transform:uppercase;margin:0 0 48px}
.meta a{color:var(--muted);text-decoration:none;border-bottom:1px solid var(--rule)}
.meta a:hover{color:var(--accent);border-color:var(--accent)}
.stats{display:grid;grid-template-columns:repeat(4,1fr);border-top:1px solid var(--ink);margin:0 0 64px}
.stats > div{padding:16px 24px 8px;border-right:1px solid var(--rule)}
.stats > div:first-child{padding-left:0}.stats > div:last-child{border-right:0}
.stats .n{font-size:48px;font-weight:400;letter-spacing:-.035em;line-height:1}
.stats .l{font-family:var(--f-mono);font-size:10px;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);margin-top:8px}
nav.toc{border-top:1px solid var(--rule);border-bottom:1px solid var(--rule);padding:16px 0;margin:0 0 64px;columns:2;column-gap:48px}
nav.toc a{display:flex;gap:12px;padding:4px 0;color:var(--ink-2);text-decoration:none;font-size:14px;break-inside:avoid}
nav.toc a span{font-family:var(--f-mono);font-size:11px;color:var(--muted-2);min-width:24px;padding-top:2px}
nav.toc a:hover{color:var(--accent)}
section{margin:0 0 64px}
section > .eyebrow{margin-bottom:8px}
h2{font-weight:500;font-size:30px;letter-spacing:-.02em;line-height:1.15;margin:0 0 8px}
.rule-accent{width:64px;height:0;border-top:2px solid var(--accent);margin:0 0 24px}
p{margin:0 0 16px;color:var(--ink-2);max-width:720px}
p strong{font-weight:600;color:var(--ink)}
ul{margin:0 0 16px;padding-left:16px;max-width:720px;color:var(--ink-2)}
li{margin:0 0 8px}
table{border-collapse:collapse;width:100%;margin:8px 0 24px;font-size:13px}
th{font-family:var(--f-mono);font-weight:500;font-size:10px;letter-spacing:.12em;text-transform:uppercase;color:var(--muted);text-align:left;padding:8px 12px 8px 0;border-bottom:1px solid var(--ink);vertical-align:bottom}
td{padding:12px 12px 12px 0;border-bottom:1px solid var(--rule);vertical-align:top;color:var(--ink-2);line-height:1.5}
td:first-child{font-family:var(--f-mono);font-size:12px;color:var(--muted);white-space:nowrap}
table.steps td:nth-child(2){font-weight:500;color:var(--ink);min-width:130px}
table.steps td:nth-child(3){min-width:260px}
a{color:var(--accent)}
.foot{font-family:var(--f-mono);font-size:10px;letter-spacing:.14em;text-transform:uppercase;color:var(--muted-2);border-top:1px solid var(--rule);padding-top:16px;margin-top:96px;display:flex;justify-content:space-between}
@media (max-width:760px){.page{padding:32px 16px 64px;margin:0}.stats{grid-template-columns:repeat(2,1fr)}nav.toc{columns:1}h1{font-size:34px}table{display:block;overflow-x:auto}}
@media print{html,body{background:#fff}.page{box-shadow:none;margin:0;padding:0}}
"""
heads=[split_head(c[1]) for c in C if c[0]=='h2']
nwf=sum(1 for h in heads if h[1])
out=[f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>BSC PLM · MTM PRD</title>{FONTS}<style>{css}</style></head><body><div class="page">']
out.append('<p class="eyebrow"><b>Product requirements</b> · Bombay Shirt Company · v1.0<span class="pill">Confidential</span></p>')
out.append('<h1>BSC PLM, the <em>made-to-measure</em> edition</h1>')
out.append(f'<p class="lede">{esc(nd(SUB))}</p>')
out.append(f'<p class="meta">{DATE} · devx Labs · <a href="bsc-plm-flow-mtm.html">View the flow diagram</a> · <a href="bsc-plm-prd-mtm-v1.0.docx">Download the Word file</a> · <a href="../">Open the prototype</a></p>')
out.append(f'<div class="stats"><div><div class="n">{nwf}</div><div class="l">Workflows</div></div><div><div class="n">6</div><div class="l">Roles</div></div><div><div class="n">5</div><div class="l">States per new fabric</div></div><div><div class="n">1</div><div class="l">Record per order</div></div></div>')
out.append('<nav class="toc">'+''.join(f'<a href="#s{i}"><span>{h[0]}</span>{esc((h[1]+": " if h[1] else "")+h[2])}</a>' for i,h in enumerate(heads))+'</nav>')
hi=-1; open_sec=False
for c in C:
    if c[0]=='h2':
        if open_sec: out.append('</section>')
        hi+=1; num,wf,title=heads[hi]; open_sec=True
        out.append(f'<section id="s{hi}"><p class="eyebrow"><b>Section {num}</b>{" / "+esc(wf) if wf else ""}</p><h2>{esc(title)}</h2><div class="rule-accent"></div>')
    elif c[0]=='p': out.append(f'<p>{para_html(c[1])}</p>')
    elif c[0]=='b': out.append(f'<p><strong>{esc(nd(c[1]))}</strong> {para_html(c[2])}</p>')
    elif c[0]=='ul': out.append('<ul>'+''.join(f'<li>{esc(nd(x))}</li>' for x in c[1])+'</ul>')
    elif c[0]=='table':
        cls=' class="steps"' if c[1][0]=='#' else ''
        out.append(f'<table{cls}><thead><tr>'+''.join(f'<th>{esc(nd(h))}</th>' for h in c[1])+'</tr></thead><tbody>')
        for r in c[2]: out.append('<tr>'+''.join(f'<td>{para_html(x)}</td>' for x in r)+'</tr>')
        out.append('</tbody></table>')
if open_sec: out.append('</section>')
out.append(f'<div class="foot"><span>BSC PLM · MTM PRD</span><span>devx Labs · Confidential</span></div></div></body></html>')
html_out='\n'.join(out)
assert '—' not in re.sub(r'<style>.*?</style>','',html_out,flags=re.S), 'em dash left in output'
open(sys.argv[1],'w').write(html_out)

# ===================================================================== DOCX · devx Doctrine
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
INK=RGBColor(0x0A,0x0A,0x0A); INK2=RGBColor(0x1A,0x1A,0x1A); MUTED=RGBColor(0x5C,0x60,0x66); ACCENT=RGBColor(0x1E,0x6F,0xFF)
def font(run, name, size=None, color=None, bold=None, italic=None):
    run.font.name=name; rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a), name)
    if size: run.font.size=Pt(size)
    if color is not None: run.font.color.rgb=color
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
def borders(el, **sides):
    pr=el._element.get_or_add_tcPr() if hasattr(el._element,'get_or_add_tcPr') else None
    b=OxmlElement('w:tcBorders')
    for side,(sz,col) in sides.items():
        e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz)); e.set(qn('w:color'),col); b.append(e)
    for side in ('top','left','bottom','right'):
        if side not in sides:
            e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'nil'); b.append(e)
    pr.append(b)
def spacing(par, before=0, after=8):
    par.paragraph_format.space_before=Pt(before); par.paragraph_format.space_after=Pt(after)
d=Document()
st=d.styles['Normal']; st.font.name='Inter Tight'; st.font.size=Pt(10); st.font.color.rgb=INK2
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Inter Tight')
for s in d.sections: s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(2.2)
pp=d.add_paragraph(); font(pp.add_run('PRODUCT REQUIREMENTS · BOMBAY SHIRT COMPANY · V1.0 · CONFIDENTIAL'),'JetBrains Mono',7.5,MUTED); spacing(pp,0,12)
pp=d.add_paragraph(); font(pp.add_run('BSC PLM, the '),'Inter Tight',28,INK,False); font(pp.add_run('made-to-measure'),'Source Serif 4',28,ACCENT,False,True); font(pp.add_run(' edition'),'Inter Tight',28,INK,False); spacing(pp,0,12)
pp=d.add_paragraph(); font(pp.add_run(nd(SUB)),'Inter Tight',11,INK2); spacing(pp,0,12)
pp=d.add_paragraph(); font(pp.add_run(f'{DATE.upper()} · DEVX LABS'),'JetBrains Mono',7.5,MUTED); spacing(pp,0,24)
def add_cell(cell, text, head=False, first=False):
    cell.text=''
    lines=nd(text).split('\n') if text else ['']
    for k,l in enumerate(lines):
        par=cell.paragraphs[0] if k==0 else cell.add_paragraph()
        spacing(par,0,2)
        if head: font(par.add_run(l.upper()),'JetBrains Mono',6.5,MUTED)
        elif first: font(par.add_run(l),'JetBrains Mono',7.5,MUTED)
        else: font(par.add_run(l),'Inter Tight',8.5,INK2)
    borders(cell, bottom=(8,'0A0A0A') if head else (4,'E5E5E5'))
hi=-1
for c in C:
    if c[0]=='h2':
        hi+=1; num,wf,title=heads[hi]
        pp=d.add_paragraph(); font(pp.add_run(f'SECTION {num}'),'JetBrains Mono',7.5,ACCENT)
        if wf: font(pp.add_run(f' / {wf.upper()}'),'JetBrains Mono',7.5,MUTED)
        spacing(pp,24,4)
        pp=d.add_paragraph(); font(pp.add_run(title),'Inter Tight',18,INK,False); spacing(pp,0,10)
    elif c[0]=='p':
        for l in c[1].split('\n'): pp=d.add_paragraph(); font(pp.add_run(nd(l)),'Inter Tight',10,INK2)
    elif c[0]=='b':
        pp=d.add_paragraph(); font(pp.add_run(nd(c[1])+' '),'Inter Tight',10,INK,True); font(pp.add_run(nd(c[2])),'Inter Tight',10,INK2)
    elif c[0]=='ul':
        for x in c[1]: pp=d.add_paragraph(style='List Bullet'); font(pp.add_run(nd(x)),'Inter Tight',10,INK2)
    elif c[0]=='table':
        t=d.add_table(rows=1, cols=len(c[1]))
        for i,h in enumerate(c[1]): add_cell(t.rows[0].cells[i], h, head=True)
        for row in c[2]:
            cells=t.add_row().cells
            for i,x in enumerate(row): add_cell(cells[i], x, first=(i==0))
        d.add_paragraph()
pp=d.add_paragraph(); font(pp.add_run('BSC PLM · MTM PRD · DEVX LABS · CONFIDENTIAL'),'JetBrains Mono',7,MUTED); spacing(pp,36,0)
d.save(sys.argv[2])
print('ok', len(C), 'blocks')
